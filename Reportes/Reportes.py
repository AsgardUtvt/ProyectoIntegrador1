import sys
import os
from PyQt6.QtWidgets import (
    QMainWindow, QWidget, QMessageBox, QTableWidget, QTableWidgetItem, QFileDialog
)
from PyQt6 import uic
from PyQt6.QtGui import QColor
from BaseDatos.MySqlManager import MySqlManager
from openpyxl import Workbook
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle
from reportlab.lib import colors

class ReportesWindow(QWidget):
    def __init__(self, db: MySqlManager, navegar=None):
        super().__init__()
        self.db = db
        self.navegar = navegar
        self.init_ui()

    def init_ui(self):
        # -------------------------
        # Carga del archivo .ui con ruta relativa segura
        # -------------------------
        directorio_actual = os.path.dirname(os.path.abspath(__file__))
        ruta_ui = os.path.join(directorio_actual, "..", "Documentacion", "QtDesigner", "Reportes_widget.ui")
        
        if not os.path.exists(ruta_ui):
            QMessageBox.critical(None, "Error crítico", f"No se encontró el archivo de interfaz en:\n{ruta_ui}")
            sys.exit(1)

        uic.loadUi(ruta_ui, self)

        # -------------------------
        # Configuración inicial de la tabla
        # -------------------------
        self.configurar_tabla()

        # -------------------------
        # Conexión de Botones y Eventos
        # -------------------------
        self.btnGenerar.clicked.connect(self.cargar_inventario)
        self.btnLimpiar.clicked.connect(self.limpiar_tabla)
        self.btnExcel.clicked.connect(self.exportar_excel)
        self.btnPDF.clicked.connect(self.exportar_pdf)

        if hasattr(self, "txtBuscar"):
            self.txtBuscar.textChanged.connect(self.buscar_medicamento)

        # Cargar los datos automáticamente al iniciar
        self.cargar_inventario()

    def configurar_tabla(self):
        """Configura las columnas y propiedades de la tabla."""
        if not hasattr(self, "tablaReporte"):
            return
        tabla = self.tablaReporte
        tabla.setColumnCount(7)
        tabla.setHorizontalHeaderLabels([
            "ID", "Medicamento", "Cantidad", "Stock Mínimo", "Caducidad", "Dosis", "Costo"
        ])
        tabla.verticalHeader().setVisible(False)
        tabla.setAlternatingRowColors(True)
        tabla.horizontalHeader().setStretchLastSection(True)
        tabla.setSelectionBehavior(tabla.SelectionBehavior.SelectRows)
        tabla.setEditTriggers(tabla.EditTrigger.NoEditTriggers)

    # --- LÓGICA DE DATOS CON MySqlManager ---

    def cargar_inventario(self):
        """Consulta la base de datos usando MySqlManager y llena la tabla."""
        try:
            query = """
                SELECT id_medicamento, medicamento_name, medicamento_cantidad, 
                       medicamento_min, medicamento_caducidad, medicamento_dosis, medicamento_costo
                FROM Medicamento
                ORDER BY medicamento_name
            """
            datos = self.db.fetchall(query)
            if datos is not None:
                self.actualizar_tabla_visual(datos)
            else:
                self.actualizar_tabla_visual([])
        except Exception as err:
            QMessageBox.warning(self, "Error de Base de Datos", f"No se pudo cargar el inventario:\n{err}")

    def buscar_medicamento(self):
        """Filtra los medicamentos en tiempo real desde el QLineEdit."""
        texto = self.txtBuscar.text().strip()
        try:
            query = """
                SELECT id_medicamento, medicamento_name, medicamento_cantidad, 
                       medicamento_min, medicamento_caducidad, medicamento_dosis, medicamento_costo
                FROM Medicamento
                WHERE medicamento_name LIKE %s
                ORDER BY medicamento_name
            """
            datos = self.db.fetchall(query, ("%" + texto + "%",))
            if datos is not None:
                self.actualizar_tabla_visual(datos)
            else:
                self.actualizar_tabla_visual([])
        except Exception as err:
            QMessageBox.critical(self, "Error", f"Error en la búsqueda:\n{err}")

    def actualizar_tabla_visual(self, datos):
        """Inserta los registros en la tabla y marca en rojo los stocks bajos."""
        tabla = self.tablaReporte
        tabla.setRowCount(len(datos))

        for fila, med in enumerate(datos):
            tabla.setItem(fila, 0, QTableWidgetItem(str(med.get("id_medicamento", ""))))
            tabla.setItem(fila, 1, QTableWidgetItem(str(med.get("medicamento_name", ""))))
            tabla.setItem(fila, 2, QTableWidgetItem(str(med.get("medicamento_cantidad", ""))))
            tabla.setItem(fila, 3, QTableWidgetItem(str(med.get("medicamento_min", ""))))
            tabla.setItem(fila, 4, QTableWidgetItem(str(med.get("medicamento_caducidad", ""))))
            tabla.setItem(fila, 5, QTableWidgetItem(str(med.get("medicamento_dosis", ""))))
            tabla.setItem(fila, 6, QTableWidgetItem(str(med.get("medicamento_costo", ""))))

            # Alerta visual si la cantidad es menor o igual al mínimo
            cantidad = med.get("medicamento_cantidad", 0)
            minimo = med.get("medicamento_min", 0)
            if cantidad <= minimo:
                for columna in range(7):
                    item = tabla.item(fila, columna)
                    if item:
                        item.setBackground(QColor(255, 210, 210))
                        item.setForeground(QColor(170, 0, 0))

        if hasattr(self, "lblTotal"):
            self.lblTotal.setText(f"Total de medicamentos: {len(datos)}")

    def limpiar_tabla(self):
        """Limpia la barra de búsqueda y vacía la tabla."""
        self.tablaReporte.setRowCount(0)
        if hasattr(self, "txtBuscar"):
            self.txtBuscar.clear()
        if hasattr(self, "lblTotal"):
            self.lblTotal.setText("Total de medicamentos: 0")

    def exportar_excel(self):
        """Exporta los datos de la tabla a un archivo Excel."""
        archivo, _ = QFileDialog.getSaveFileName(self, "Guardar reporte Excel", "", "Excel (*.xlsx)")
        if not archivo:
            return

        libro = Workbook()
        hoja = libro.active
        hoja.title = "Inventario"

        for col in range(self.tablaReporte.columnCount()):
            hoja.cell(row=1, column=col + 1).value = self.tablaReporte.horizontalHeaderItem(col).text()

        for fila in range(self.tablaReporte.rowCount()):
            for col in range(self.tablaReporte.columnCount()):
                item = self.tablaReporte.item(fila, col)
                if item:
                    hoja.cell(row=fila + 2, column=col + 1).value = item.text()

        libro.save(archivo)
        QMessageBox.information(self, "Éxito", "Reporte exportado a Excel correctamente.")

    def exportar_pdf(self):
        """Exporta los datos de la tabla a un documento PDF."""
        archivo, _ = QFileDialog.getSaveFileName(self, "Guardar PDF", "", "PDF (*.pdf)")
        if not archivo:
            return

        datos_tabla = []
        encabezados = []

        for c in range(self.tablaReporte.columnCount()):
            encabezados.append(self.tablaReporte.horizontalHeaderItem(c).text())
        datos_tabla.append(encabezados)

        for fila in range(self.tablaReporte.rowCount()):
            registro = []
            for col in range(self.tablaReporte.columnCount()):
                item = self.tablaReporte.item(fila, col)
                registro.append(item.text() if item else "")
            datos_tabla.append(registro)

        pdf = SimpleDocTemplate(archivo)
        tabla = Table(datos_tabla)
        tabla.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1565C0")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 1, colors.black),
            ("BACKGROUND", (0, 1), (-1, -1), colors.beige),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("BOTTOMPADDING", (0, 0), (-1, 0), 10)
        ]))

        pdf.build([tabla])
        QMessageBox.information(self, "Éxito", "PDF generado correctamente.")